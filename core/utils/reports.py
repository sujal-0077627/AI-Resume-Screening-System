"""Generate beautifully formatted reports (Excel/CSV) using Pandas and openpyxl."""
import io
import pandas as pd
from django.http import HttpResponse
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def _build_dataframe(scores):
    """Build a pandas DataFrame from Score queryset."""
    rows = []
    for score in scores:
        rows.append({
            'Candidate Name': score.candidate.name or f"Candidate #{score.candidate.id}",
            'Email': score.candidate.email or "-",
            'Phone': score.candidate.phone or "-",
            'Job Title': score.job.title,
            'Overall Score': f"{score.overall_score:.1f}%",
            'TF-IDF Score': f"{score.tfidf_score:.1f}%",
            'Semantic Score': f"{score.semantic_score:.1f}%",
            'Matched Skills': score.matched_skills or "-",
            'Missing Skills': score.missing_skills or "-",
            'Status': score.get_status_display(),
            'Email Sent': 'Yes' if score.email_sent else 'No',
            'Screened At': score.created_at.strftime('%Y-%m-%d %H:%M'),
        })
    return pd.DataFrame(rows)


def export_excel(scores, filename='resume_screening_report.xlsx'):
    """Export scores to a beautifully formatted Excel file and return an HttpResponse."""
    df = _build_dataframe(scores)
    buffer = io.BytesIO()
    
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Screening Results')
        ws = writer.sheets['Screening Results']

        # Colors & Styling
        HEADER_FILL = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid") # Deep Indigo
        HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid") # Light Slate
        WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        
        # Status Fills & Fonts
        SELECTED_FILL = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        SELECTED_FONT = Font(name="Calibri", size=11, bold=True, color="166534")
        
        REJECTED_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        REJECTED_FONT = Font(name="Calibri", size=11, bold=True, color="991B1B")

        PENDING_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
        PENDING_FONT = Font(name="Calibri", size=11, bold=True, color="92400E")

        # Borders
        THIN_BORDER = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        # Style Header Row
        ws.row_dimensions[1].height = 26
        for col_idx in range(1, len(df.columns) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Style Data Rows
        for row_idx in range(2, len(df) + 2):
            ws.row_dimensions[row_idx].height = 24
            is_zebra = (row_idx % 2 == 0)
            
            for col_idx, col_name in enumerate(df.columns, start=1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = THIN_BORDER
                cell.font = Font(name="Calibri", size=11, color="1E293B")
                
                # Base Fill
                cell.fill = ZEBRA_FILL if is_zebra else WHITE_FILL
                
                # Alignment
                if col_name in ('Candidate Name', 'Email', 'Job Title', 'Matched Skills', 'Missing Skills'):
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                # Highlight Status column
                if col_name == 'Status':
                    val = str(cell.value or '').lower()
                    if 'select' in val:
                        cell.fill = SELECTED_FILL
                        cell.font = SELECTED_FONT
                    elif 'reject' in val:
                        cell.fill = REJECTED_FILL
                        cell.font = REJECTED_FONT
                    else:
                        cell.fill = PENDING_FILL
                        cell.font = PENDING_FONT

        # Auto-fit Column Widths cleanly
        for col_idx, col_name in enumerate(df.columns, start=1):
            col_letter = get_column_letter(col_idx)
            # Find max string length in column
            max_len = max(len(str(col_name)), df[col_name].astype(str).map(len).max() if not df.empty else 0)
            
            # Width bounds
            if col_name in ('Matched Skills', 'Missing Skills'):
                width = min(45, max(25, max_len + 3))
            elif col_name in ('Candidate Name', 'Job Title'):
                width = min(30, max(18, max_len + 4))
            elif col_name == 'Email':
                width = min(35, max(20, max_len + 4))
            else:
                width = max(14, max_len + 4)
                
            ws.column_dimensions[col_letter].width = width

    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def export_csv(scores, filename='resume_screening_report.csv'):
    """Export scores to CSV with UTF-8 BOM so Excel opens with clean aligned columns."""
    df = _build_dataframe(scores)
    buffer = io.StringIO()
    df.to_csv(buffer, index=False)
    buffer.seek(0)

    response = HttpResponse(buffer.getvalue().encode('utf-8-sig'), content_type='text/csv; charset=utf-8-sig')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response